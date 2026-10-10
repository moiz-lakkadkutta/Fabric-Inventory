"""#216 items 1 + 2 — credit-note series and the credit-note document.

1. Numbering (Moiz, locked): a §34 credit note gets its own per-firm series
   ``CN/<FY>`` where FY is the Indian financial year (1 Apr - 31 Mar) of the
   note's date, e.g. ``CN/26-27/0001``. The counter is gapless per (org, firm,
   CREDIT_NOTE, series) so it restarts every 1 April. ≤16 chars (Rule 46(b) /
   53(1A)(b)). A SAME-PERIOD cancellation is a true cancellation, not a credit
   note (#199 / #214 CA sheet item 4), so its GL reversal must NOT consume a
   CN serial (otherwise the CN series reported in GSTR-1 would have gaps); it
   gets an internal ``CXL/<FY>`` series instead. Existing vouchers keep their
   numbers (no data rewrite).
2. Document: ``GET /invoices/{id}/credit-note/pdf`` renders a Rule 53(1A)
   credit note for a cross-period cancel; 404 when the invoice has no §34
   credit note (not cancelled, or cancelled in the same period).

GATED — GST numbering/document: PENDING MOIZ + CA SIGN-OFF.
"""

from __future__ import annotations

import datetime
import io
import uuid
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.engine import Engine
from sqlalchemy.orm import Session as OrmSession

from app.exceptions import NotFoundError
from app.models import Firm, SalesInvoice, Voucher
from app.models.accounting import VoucherType
from app.service import gst_service, pdf_service, reports_service, sales_service
from tests.test_concurrency_postings import (
    _drop_org,
    _new_session,
    _race,
    _seed_org_firm_party_item,
)
from tests.test_invoice_pdf_routers import _seed_party_and_item
from tests.test_reports_gstr1 import _seed_gstr1_recon_org
from tests.test_reports_routers import _auth, _signup_owner
from tests.test_sales_invoice_cancel_gst_period import _cancel, _credit_note, _invoice, _party

_UTC = datetime.UTC
_GSTIN = "27ABCDE1234F1Z5"
_OCT10 = datetime.datetime(2026, 10, 10, 6, tzinfo=_UTC)


# ──────────────────────────────────────────────────────────────────────
# 1. Series label
# ──────────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("d", "series"),
    [
        (datetime.date(2026, 10, 10), "CN/26-27"),
        (datetime.date(2026, 4, 1), "CN/26-27"),
        (datetime.date(2027, 3, 31), "CN/26-27"),
        (datetime.date(2027, 4, 1), "CN/27-28"),
        (datetime.date(2099, 6, 1), "CN/99-00"),
    ],
)
def test_credit_note_series_is_cn_slash_indian_fy(d: datetime.date, series: str) -> None:
    assert gst_service.credit_note_series(d) == series


def test_credit_note_full_number_fits_16_chars() -> None:
    # Rule 46(b) / 53(1A)(b): serial number not exceeding 16 characters.
    full = f"{gst_service.credit_note_series(datetime.date(2026, 10, 10))}/9999"
    assert full == "CN/26-27/9999"
    assert len(full) <= 16


# ──────────────────────────────────────────────────────────────────────
# 2. Allocation on cancel
# ──────────────────────────────────────────────────────────────────────


def _setup(
    db_session: OrmSession, *, gstin: str | None = _GSTIN
) -> tuple[uuid.UUID, uuid.UUID, uuid.UUID, uuid.UUID]:
    org_id, firm_id, item_id = _seed_gstr1_recon_org(db_session)
    party_id = _party(db_session, org_id, state_code="MH", gstin=gstin)
    return org_id, firm_id, item_id, party_id


def test_cross_period_cancel_gets_cn_series_not_invoice_series(db_session: OrmSession) -> None:
    org_id, firm_id, item_id, party_id = _setup(db_session)
    inv = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2026, 9, 25),
    )
    _cancel(db_session, org_id=org_id, invoice=inv, at=_OCT10)
    cn = _credit_note(db_session, org_id=org_id, invoice=inv)
    assert (cn.series, cn.number) == ("CN/26-27", "0001")
    assert cn.series != inv.series
    assert len(f"{cn.series}/{cn.number}") <= 16


def test_cn_counter_is_gapless_and_resets_each_fy(db_session: OrmSession) -> None:
    org_id, firm_id, item_id, party_id = _setup(db_session)

    def inv_on(d: datetime.date) -> SalesInvoice:
        return _invoice(
            db_session,
            org_id=org_id,
            firm_id=firm_id,
            party_id=party_id,
            item_id=item_id,
            invoice_date=d,
        )

    a = inv_on(datetime.date(2025, 1, 10))
    b = inv_on(datetime.date(2025, 1, 11))
    c = inv_on(datetime.date(2025, 3, 5))
    _cancel(db_session, org_id=org_id, invoice=a, at=datetime.datetime(2025, 2, 10, 6, tzinfo=_UTC))
    _cancel(db_session, org_id=org_id, invoice=b, at=datetime.datetime(2025, 3, 31, 6, tzinfo=_UTC))
    # 1 April → new FY → counter restarts.
    _cancel(db_session, org_id=org_id, invoice=c, at=datetime.datetime(2025, 4, 1, 6, tzinfo=_UTC))

    got = [
        (n.series, n.number)
        for n in (_credit_note(db_session, org_id=org_id, invoice=x) for x in (a, b, c))
    ]
    assert got == [("CN/24-25", "0001"), ("CN/24-25", "0002"), ("CN/25-26", "0001")]


def test_same_period_cancel_does_not_consume_a_cn_serial(db_session: OrmSession) -> None:
    """A same-month cancel is a true cancellation (no §34 note, not in
    GSTR-1). Its GL reversal gets the internal CXL/<FY> series so the CN
    series reported in CDNR/CDNUR stays consecutive."""
    org_id, firm_id, item_id, party_id = _setup(db_session)
    same = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2026, 10, 2),
    )
    cross = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2026, 9, 20),
    )
    at = datetime.datetime(2026, 10, 10, 6, tzinfo=_UTC)
    _cancel(db_session, org_id=org_id, invoice=same, at=at)
    _cancel(db_session, org_id=org_id, invoice=cross, at=at)

    rev = _credit_note(db_session, org_id=org_id, invoice=same)
    cn = _credit_note(db_session, org_id=org_id, invoice=cross)
    assert (rev.series, rev.number) == ("CXL/26-27", "0001")
    assert (cn.series, cn.number) == ("CN/26-27", "0001")


def test_legacy_credit_note_keeps_its_number_and_new_series_is_independent(
    db_session: OrmSession,
) -> None:
    """A pre-#216 CREDIT_NOTE in the invoice series ("RT/2526"/"0001") is
    left alone; the first new note is still CN/<FY>/0001."""
    org_id, firm_id, item_id, party_id = _setup(db_session)
    inv = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2026, 9, 20),
    )
    legacy = Voucher(
        org_id=org_id,
        firm_id=firm_id,
        voucher_type=VoucherType.CREDIT_NOTE,
        series=inv.series,
        number="0001",
        voucher_date=datetime.date(2026, 9, 1),
        narration="legacy note",
        total_debit=Decimal("0"),
        total_credit=Decimal("0"),
    )
    db_session.add(legacy)
    db_session.flush()
    _cancel(db_session, org_id=org_id, invoice=inv, at=_OCT10)
    cn = _credit_note(db_session, org_id=org_id, invoice=inv)
    assert (cn.series, cn.number) == ("CN/26-27", "0001")
    db_session.refresh(legacy)
    assert (legacy.series, legacy.number) == (inv.series, "0001")


# ──────────────────────────────────────────────────────────────────────
# 3. GSTR-1 rows + xlsx carry the full note number
# ──────────────────────────────────────────────────────────────────────


def test_gstr1_cdnr_row_carries_full_cn_number(db_session: OrmSession) -> None:
    org_id, firm_id, item_id, party_id = _setup(db_session)
    inv = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2026, 9, 25),
    )
    _cancel(db_session, org_id=org_id, invoice=inv, at=_OCT10)
    oct_ = reports_service.compute_gstr1(
        db_session, org_id=org_id, firm_id=firm_id, period="2026-10"
    )
    assert len(oct_.cdnr) == 1
    note = oct_.cdnr[0]
    assert (note.note_series, note.note_number) == ("CN/26-27", "0001")
    assert note.note_no == "CN/26-27/0001"
    assert note.invoice_no == f"{inv.series}/{inv.number}"


def test_gstr1_xlsx_note_column_is_full_number(
    http_client: TestClient, sync_engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    from tests.test_reports_gstr1 import _seed_b2b_party, _seed_item
    from tests.test_reports_routers import _create_and_finalize_invoice

    me = _signup_owner(http_client)
    org_id = uuid.UUID(me["org_id"])
    party_id = _seed_b2b_party(sync_engine, org_id=org_id)
    item_id = _seed_item(sync_engine, org_id=org_id)
    invoice_id = _create_and_finalize_invoice(
        http_client, me, party_id=party_id, item_id=item_id, invoice_date="2026-09-25"
    )
    monkeypatch.setattr(
        sales_service, "_utcnow", lambda: datetime.datetime(2026, 10, 10, 6, tzinfo=_UTC)
    )
    resp = http_client.post(
        f"/invoices/{invoice_id}/cancel",
        headers=_auth(me["access_token"]),
        json={"reason": "goods rejected"},
    )
    assert resp.status_code == 200, resp.text
    inv = resp.json()

    js = http_client.get("/reports/gstr1?period=2026-10", headers=_auth(me["access_token"]))
    assert js.status_code == 200, js.text
    note = js.json()["cdnr"][0]
    assert note["note_no"] == "CN/26-27/0001"

    xlsx = http_client.get(
        "/reports/gstr1?period=2026-10&format=xlsx", headers=_auth(me["access_token"])
    )
    assert xlsx.status_code == 200, xlsx.text
    ws = load_workbook(io.BytesIO(xlsx.content))["CDNR"]
    row = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]], strict=True))
    assert row["Note #"] == "CN/26-27/0001"
    assert row["Original invoice #"] == f"{inv['series']}/{inv['number']}"


# ──────────────────────────────────────────────────────────────────────
# 4. Concurrency — two cancels at once get distinct CN numbers
# ──────────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("n", [2, 4])
def test_concurrent_cancels_get_distinct_cn_numbers(
    sync_engine: Engine, admin_engine: Engine, n: int
) -> None:
    org_id, firm_id, party_id, item_id = _seed_org_firm_party_item(sync_engine)
    try:
        inv_ids: list[uuid.UUID] = []
        with _new_session(sync_engine, org_id) as s:
            for _ in range(n):
                inv = sales_service.create_draft_invoice(
                    s,
                    org_id=org_id,
                    firm_id=firm_id,
                    party_id=party_id,
                    invoice_date=datetime.date(2026, 9, 15),
                    ship_to_state="MH",
                    lines=[{"item_id": item_id, "qty": "1", "price": "1000", "gst_rate": "5"}],
                )
                sales_service.finalize_invoice(
                    s, org_id=org_id, sales_invoice_id=inv.sales_invoice_id
                )
                inv_ids.append(inv.sales_invoice_id)
            s.commit()

        queue = list(inv_ids)
        import threading

        qlock = threading.Lock()

        def cancel_one(s: OrmSession) -> object:
            with qlock:
                target = queue.pop()
            return sales_service.cancel_invoice(
                s,
                org_id=org_id,
                sales_invoice_id=target,
                reason="race",
                now=datetime.datetime(2026, 10, 10, 6, tzinfo=_UTC),
            )

        results = _race(sync_engine, org_id, n, cancel_one)
        assert results == ["OK"] * n, results

        with _new_session(sync_engine, org_id) as s:
            numbers = (
                s.execute(
                    text(
                        "SELECT series || '/' || number FROM voucher WHERE org_id = :o "
                        "AND voucher_type = 'CREDIT_NOTE' AND deleted_at IS NULL ORDER BY number"
                    ),
                    {"o": str(org_id)},
                )
                .scalars()
                .all()
            )
        assert numbers == [f"CN/26-27/{i:04d}" for i in range(1, n + 1)]
    finally:
        _drop_org(admin_engine, org_id)


# ──────────────────────────────────────────────────────────────────────
# 5. Credit-note document (Rule 53(1A))
# ──────────────────────────────────────────────────────────────────────


def test_render_credit_note_html_contains_rule_53_1a_fields(db_session: OrmSession) -> None:
    org_id, firm_id, item_id, party_id = _setup(db_session)
    firm = db_session.execute(select(Firm).where(Firm.firm_id == firm_id)).scalar_one()
    firm.address = "142, Resham Bhavan, Ring Road, Surat 395002"
    firm.legal_name = "Rajesh Textiles Pvt Ltd"
    from app.utils.crypto import encrypt_pii, get_org_dek

    firm.gstin = encrypt_pii(
        "27AAACR1234A1Z1", dek=get_org_dek(db_session, org_id=org_id), org_id=org_id
    )
    db_session.flush()
    inv = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2026, 9, 25),
        qty="10",
        price="1000",
        gst_rate="5",
    )
    _cancel(db_session, org_id=org_id, invoice=inv, at=_OCT10)

    html = pdf_service.render_credit_note_html(
        db_session, invoice_id=inv.sales_invoice_id, org_id=org_id
    )
    # (a) supplier name, address, GSTIN
    assert "Rajesh Textiles Pvt Ltd" in html
    assert "Resham Bhavan" in html
    assert "27AAACR1234A1Z1" in html
    # (b) nature of document
    assert "Credit Note" in html
    # (c) consecutive serial ≤16 chars, (d) date of issue
    assert "CN/26-27/0001" in html
    assert "10 Oct 2026" in html
    # (e) recipient name + GSTIN (registered)
    assert _GSTIN in html
    # (g) serial number + date of the corresponding tax invoice
    assert f"{inv.series}/{inv.number}" in html
    assert "25 Sep 2026" in html
    # (h) value, taxable value, rate, tax credited — CGST/SGST split per line
    assert "10,000.00" in html  # taxable
    assert "250.00" in html  # CGST == SGST
    assert "10,500.00" in html  # total credited
    assert "5.00" in html  # rate
    assert "5208" in html  # HSN
    # (i) signature of supplier / authorised signatory
    assert "Authorised Signatory" in html
    assert "For Rajesh Textiles Pvt Ltd" in html
    # Reason is recorded on the note.
    assert "customer rejected" in html


def test_credit_note_unregistered_over_50k_shows_delivery_place_and_state(
    db_session: OrmSession,
) -> None:
    org_id, firm_id, item_id, party_id = _setup(db_session, gstin=None)
    inv = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2026, 9, 25),
        qty="1",
        price="60000",
        gst_rate="5",
    )
    inv.bill_to_address = "12 Linking Road, Mumbai"
    inv.ship_to_address = "Godown 7, Bhiwandi"
    db_session.flush()
    _cancel(db_session, org_id=org_id, invoice=inv, at=_OCT10)
    html = pdf_service.render_credit_note_html(
        db_session, invoice_id=inv.sales_invoice_id, org_id=org_id
    )
    assert "Unregistered" in html
    assert "12 Linking Road, Mumbai" in html
    assert "Place of Delivery" in html and "Godown 7, Bhiwandi" in html
    assert "Maharashtra" in html


def test_credit_note_html_404_for_same_period_cancel(db_session: OrmSession) -> None:
    org_id, firm_id, item_id, party_id = _setup(db_session)
    inv = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2026, 10, 2),
    )
    _cancel(db_session, org_id=org_id, invoice=inv, at=_OCT10)
    with pytest.raises(NotFoundError):
        pdf_service.render_credit_note_html(
            db_session, invoice_id=inv.sales_invoice_id, org_id=org_id
        )


def _cancel_via_api(
    http_client: TestClient,
    me: dict[str, str],
    invoice_id: str,
    monkeypatch: pytest.MonkeyPatch,
    at: datetime.datetime,
) -> dict[str, object]:
    monkeypatch.setattr(sales_service, "_utcnow", lambda: at)
    resp = http_client.post(
        f"/invoices/{invoice_id}/cancel",
        headers=_auth(me["access_token"]),
        json={"reason": "goods rejected"},
    )
    assert resp.status_code == 200, resp.text
    body: dict[str, object] = resp.json()
    return body


def _api_invoice(http_client: TestClient, sync_engine: Engine, me: dict[str, str]) -> str:
    from tests.test_invoice_pdf_routers import _create_and_finalize

    party_id, item_id = _seed_party_and_item(
        sync_engine, org_id=uuid.UUID(me["org_id"]), firm_id=uuid.UUID(me["firm_id"])
    )
    return _create_and_finalize(http_client, me, party_id, item_id)  # dated 2026-04-30


def test_credit_note_pdf_endpoint_returns_pdf_and_detail_exposes_note(
    http_client: TestClient, sync_engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    me = _signup_owner(http_client)
    invoice_id = _api_invoice(http_client, sync_engine, me)

    # Before cancel: no credit note → 404, and the detail carries none.
    resp = http_client.get(
        f"/invoices/{invoice_id}/credit-note/pdf", headers=_auth(me["access_token"])
    )
    assert resp.status_code == 404, resp.text
    detail = http_client.get(f"/invoices/{invoice_id}", headers=_auth(me["access_token"]))
    assert detail.json()["credit_note"] is None

    cancelled = _cancel_via_api(
        http_client, me, invoice_id, monkeypatch, datetime.datetime(2026, 5, 10, 6, tzinfo=_UTC)
    )
    note = cancelled["credit_note"]
    assert isinstance(note, dict)
    assert note["number"] == "CN/26-27/0001"
    assert note["note_date"] == "2026-05-10"
    assert uuid.UUID(note["voucher_id"])

    detail = http_client.get(f"/invoices/{invoice_id}", headers=_auth(me["access_token"]))
    assert detail.json()["credit_note"]["number"] == "CN/26-27/0001"

    resp = http_client.get(
        f"/invoices/{invoice_id}/credit-note/pdf", headers=_auth(me["access_token"])
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content.startswith(b"%PDF")
    assert "CN_26-27_0001.pdf" in resp.headers["content-disposition"]


def test_credit_note_pdf_404_for_same_period_cancel(
    http_client: TestClient, sync_engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    me = _signup_owner(http_client)
    invoice_id = _api_invoice(http_client, sync_engine, me)
    cancelled = _cancel_via_api(
        http_client, me, invoice_id, monkeypatch, datetime.datetime(2026, 4, 30, 8, tzinfo=_UTC)
    )
    assert cancelled["credit_note"] is None
    resp = http_client.get(
        f"/invoices/{invoice_id}/credit-note/pdf", headers=_auth(me["access_token"])
    )
    assert resp.status_code == 404, resp.text


def test_credit_note_pdf_cross_org_returns_404(
    http_client: TestClient, sync_engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    me_a = _signup_owner(http_client)
    invoice_id = _api_invoice(http_client, sync_engine, me_a)
    _cancel_via_api(
        http_client, me_a, invoice_id, monkeypatch, datetime.datetime(2026, 5, 10, 6, tzinfo=_UTC)
    )
    me_b = _signup_owner(http_client)
    resp = http_client.get(
        f"/invoices/{invoice_id}/credit-note/pdf", headers=_auth(me_b["access_token"])
    )
    assert resp.status_code == 404, resp.text
    # Sanity: the owner can fetch it.
    ok = http_client.get(
        f"/invoices/{invoice_id}/credit-note/pdf", headers=_auth(me_a["access_token"])
    )
    assert ok.status_code == 200, ok.text


def test_openapi_includes_credit_note_pdf_endpoint(http_client: TestClient) -> None:
    spec = http_client.get("/openapi.json").json()
    paths = [p for p in spec["paths"] if p.endswith("/invoices/{sales_invoice_id}/credit-note/pdf")]
    assert paths, list(spec["paths"])[:10]
    op = spec["paths"][paths[0]]["get"]
    assert "application/pdf" in op["responses"]["200"]["content"]
