"""#199 CA-review correction — GST-period semantics of cancelling a finalized
sales invoice (CGST Act §34 + GSTN GSTR-1 reporting).

Periods are calendar months in Asia/Kolkata.

  * Same-period cancel (cancel month IST == invoice month): a true
    cancellation before the period's GSTR-1 could be filed — the invoice is
    excluded from GSTR-1 and the reversal is dated in the same month.
  * Cross-period cancel: a §34 credit note. The original month's GSTR-1 still
    carries the invoice; the cancel month reports the credit note in CDNR
    (B2B original), CDNUR (B2CL / export original) or nets it off B2CS (B2CS
    original) and reduces the HSN summary.
  * §34 time limit: no credit note after 30-Nov following the end of the
    original invoice's financial year.

Books == return is checked per period: Σ GSTR-1 tax (credit notes negative)
== net CR movement on ledger 2100 from SALES_INVOICE + CREDIT_NOTE vouchers.
"""

from __future__ import annotations

import datetime
import io
import uuid
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.engine import Engine
from sqlalchemy.orm import Session as OrmSession

from app.exceptions import InvoiceStateError
from app.models import Ledger, Party, SalesInvoice, Voucher, VoucherLine
from app.models.accounting import JournalLineType, VoucherType
from app.models.sales import InvoiceLifecycleStatus
from app.service import reports_service, sales_service
from app.utils.crypto import encrypt_pii, get_org_dek
from tests.test_reports_gstr1 import _seed_b2b_party, _seed_gstr1_recon_org, _seed_item
from tests.test_reports_routers import _auth, _create_and_finalize_invoice, _signup_owner

_UTC = datetime.UTC
_GSTIN = "27ABCDE1234F1Z5"


# ──────────────────────────────────────────────────────────────────────
# Helpers
# ──────────────────────────────────────────────────────────────────────


def _party(
    session: OrmSession,
    org_id: uuid.UUID,
    *,
    state_code: str | None,
    gstin: str | None = None,
    is_export: bool = False,
) -> uuid.UUID:
    enc = None
    if gstin is not None:
        enc = encrypt_pii(gstin, dek=get_org_dek(session, org_id=org_id), org_id=org_id)
    party = Party(
        org_id=org_id,
        code=f"P{uuid.uuid4().hex[:6].upper()}",
        name=f"Party {uuid.uuid4().hex[:4]}",
        is_customer=True,
        state_code=state_code,
        gstin=enc,
        is_export=is_export,
    )
    session.add(party)
    session.flush()
    return party.party_id


def _invoice(
    session: OrmSession,
    *,
    org_id: uuid.UUID,
    firm_id: uuid.UUID,
    party_id: uuid.UUID,
    item_id: uuid.UUID,
    invoice_date: datetime.date,
    qty: str = "1",
    price: str = "1000",
    gst_rate: str = "5",
    ship_to_state: str | None = None,
) -> SalesInvoice:
    inv = sales_service.create_draft_invoice(
        session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        invoice_date=invoice_date,
        ship_to_state=ship_to_state,
        lines=[
            {
                "item_id": item_id,
                "qty": Decimal(qty),
                "price": Decimal(price),
                "gst_rate": Decimal(gst_rate),
                "sequence": 1,
            }
        ],
    )
    sales_service.finalize_invoice(session, org_id=org_id, sales_invoice_id=inv.sales_invoice_id)
    return inv


def _cancel(
    session: OrmSession, *, org_id: uuid.UUID, invoice: SalesInvoice, at: datetime.datetime
) -> SalesInvoice:
    return sales_service.cancel_invoice(
        session,
        org_id=org_id,
        sales_invoice_id=invoice.sales_invoice_id,
        reason="customer rejected",
        now=at,
    )


def _credit_note(session: OrmSession, *, org_id: uuid.UUID, invoice: SalesInvoice) -> Voucher:
    orig = session.execute(
        select(Voucher).where(
            Voucher.org_id == org_id,
            Voucher.voucher_type == VoucherType.SALES_INVOICE,
            Voucher.reference_id == invoice.sales_invoice_id,
        )
    ).scalar_one()
    return session.execute(
        select(Voucher).where(
            Voucher.org_id == org_id,
            Voucher.voucher_type == VoucherType.CREDIT_NOTE,
            Voucher.reference_id == orig.voucher_id,
        )
    ).scalar_one()


def _period_bounds(period: str) -> tuple[datetime.date, datetime.date]:
    y, m = (int(p) for p in period.split("-"))
    start = datetime.date(y, m, 1)
    nxt = datetime.date(y + (m == 12), m % 12 + 1, 1)
    return start, nxt - datetime.timedelta(days=1)


def _gl_2100_net(
    session: OrmSession, *, org_id: uuid.UUID, firm_id: uuid.UUID, period: str
) -> Decimal:
    """Net CR (CR - DR) on ledger 2100 from SALES_INVOICE + CREDIT_NOTE
    vouchers dated in ``period``."""
    start, end = _period_bounds(period)
    ledger_id = session.execute(
        select(Ledger.ledger_id).where(Ledger.org_id == org_id, Ledger.code == "2100")
    ).scalar_one()
    net = Decimal("0")
    rows = session.execute(
        select(VoucherLine.line_type, VoucherLine.amount)
        .join(Voucher, Voucher.voucher_id == VoucherLine.voucher_id)
        .where(
            VoucherLine.ledger_id == ledger_id,
            Voucher.firm_id == firm_id,
            Voucher.deleted_at.is_(None),
            Voucher.voucher_type.in_([VoucherType.SALES_INVOICE, VoucherType.CREDIT_NOTE]),
            Voucher.voucher_date >= start,
            Voucher.voucher_date <= end,
        )
    ).all()
    for line_type, amount in rows:
        net += Decimal(amount) if line_type == JournalLineType.CR else -Decimal(amount)
    return net


def _gstr1_net_tax(result: reports_service._Gstr1Result) -> Decimal:
    """Σ tax over all GSTR-1 supply sections, credit notes negative. B2CS is
    already netted (negative amounts) so it is summed as-is."""
    supplies = sum(
        (
            r.cgst + r.sgst + r.igst
            for bucket in (result.b2b, result.b2cl, result.b2cs, result.export)
            for r in bucket
        ),
        Decimal("0"),
    )
    notes = sum(
        (r.cgst + r.sgst + r.igst for bucket in (result.cdnr, result.cdnur) for r in bucket),
        Decimal("0"),
    )
    return supplies - notes


def _assert_books_equal_return(
    session: OrmSession, *, org_id: uuid.UUID, firm_id: uuid.UUID, period: str
) -> reports_service._Gstr1Result:
    result = reports_service.compute_gstr1(session, org_id=org_id, firm_id=firm_id, period=period)
    gl = _gl_2100_net(session, org_id=org_id, firm_id=firm_id, period=period)
    ret = _gstr1_net_tax(result)
    assert gl == ret, f"{period}: books (2100 net {gl}) != return (GSTR-1 net tax {ret})"
    return result


# ──────────────────────────────────────────────────────────────────────
# 1. Same-period cancel
# ──────────────────────────────────────────────────────────────────────


def test_same_month_cancel_excluded_from_gstr1_and_books_equal_return(
    db_session: OrmSession,
) -> None:
    org_id, firm_id, item_id = _seed_gstr1_recon_org(db_session)
    party_id = _party(db_session, org_id, state_code="MH", gstin=_GSTIN)
    inv = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2024, 9, 5),
    )
    _cancel(
        db_session, org_id=org_id, invoice=inv, at=datetime.datetime(2024, 9, 20, 6, tzinfo=_UTC)
    )

    cn = _credit_note(db_session, org_id=org_id, invoice=inv)
    assert cn.voucher_date == datetime.date(2024, 9, 20)

    result = _assert_books_equal_return(
        db_session, org_id=org_id, firm_id=firm_id, period="2024-09"
    )
    assert result.b2b == [] and result.cdnr == [] and result.cdnur == []
    assert result.hsn == []


def test_same_period_cancel_reversal_dated_in_ist(db_session: OrmSession) -> None:
    """Invoice dated 01-Oct; cancel at 2024-09-30T19:00Z = 01-Oct 00:30 IST.
    In UTC that is September (a month BEFORE the invoice); in IST it is the
    same period. The reversal must be dated 01-Oct (IST) → same-period
    cancel, excluded from October GSTR-1, and September untouched."""
    org_id, firm_id, item_id = _seed_gstr1_recon_org(db_session)
    party_id = _party(db_session, org_id, state_code="MH", gstin=_GSTIN)
    inv = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2024, 10, 1),
    )
    _cancel(
        db_session, org_id=org_id, invoice=inv, at=datetime.datetime(2024, 9, 30, 19, tzinfo=_UTC)
    )
    cn = _credit_note(db_session, org_id=org_id, invoice=inv)
    assert cn.voucher_date == datetime.date(2024, 10, 1)

    oct_ = _assert_books_equal_return(db_session, org_id=org_id, firm_id=firm_id, period="2024-10")
    assert oct_.b2b == [] and oct_.cdnr == []
    sep = _assert_books_equal_return(db_session, org_id=org_id, firm_id=firm_id, period="2024-09")
    assert sep.b2b == [] and sep.cdnr == []


# ──────────────────────────────────────────────────────────────────────
# 2. Cross-period B2B → CDNR
# ──────────────────────────────────────────────────────────────────────


def test_cross_period_b2b_cancel_reports_cdnr_in_cancel_month(db_session: OrmSession) -> None:
    org_id, firm_id, item_id = _seed_gstr1_recon_org(db_session)
    party_id = _party(db_session, org_id, state_code="MH", gstin=_GSTIN)
    inv = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2024, 8, 25),
    )
    cancelled = _cancel(
        db_session, org_id=org_id, invoice=inv, at=datetime.datetime(2024, 9, 10, 6, tzinfo=_UTC)
    )
    assert cancelled.lifecycle_status == InvoiceLifecycleStatus.CANCELLED
    cn = _credit_note(db_session, org_id=org_id, invoice=inv)
    assert cn.voucher_date == datetime.date(2024, 9, 10)

    # August: original invoice still reported in B2B, normal amounts.
    aug = _assert_books_equal_return(db_session, org_id=org_id, firm_id=firm_id, period="2024-08")
    assert len(aug.b2b) == 1
    row = aug.b2b[0]
    assert row.sales_invoice_id == inv.sales_invoice_id
    assert (row.taxable_value, row.cgst, row.sgst, row.igst) == (
        Decimal("1000.00"),
        Decimal("25.00"),
        Decimal("25.00"),
        Decimal("0"),
    )
    assert aug.cdnr == [] and aug.cdnur == []
    assert len(aug.hsn) == 1 and aug.hsn[0].taxable_value == Decimal("1000.00")

    # September: one CDNR row for the credit note; no B2B.
    sep = _assert_books_equal_return(db_session, org_id=org_id, firm_id=firm_id, period="2024-09")
    assert sep.b2b == [] and sep.cdnur == []
    assert len(sep.cdnr) == 1
    note = sep.cdnr[0]
    assert note.note_type == "C"
    assert note.note_series == cn.series and note.note_number == cn.number
    assert note.note_date == datetime.date(2024, 9, 10)
    assert note.sales_invoice_id == inv.sales_invoice_id
    assert note.invoice_number == inv.number and note.invoice_series == inv.series
    assert note.invoice_date == datetime.date(2024, 8, 25)
    assert note.gstin == _GSTIN
    assert note.place_of_supply_state == "MH"
    assert note.gst_rate == Decimal("5")
    assert note.note_value == Decimal("1050.00")
    assert (note.taxable_value, note.cgst, note.sgst, note.igst) == (
        Decimal("1000.00"),
        Decimal("25.00"),
        Decimal("25.00"),
        Decimal("0"),
    )
    # HSN reduced by the credit note's lines (net-of-CDN, GSTR-1 Table 12).
    assert len(sep.hsn) == 1
    assert sep.hsn[0].taxable_value == Decimal("-1000.00")
    assert sep.hsn[0].total_qty == Decimal("-1")


def test_cdnr_gstin_masked_without_pii(db_session: OrmSession) -> None:
    org_id, firm_id, item_id = _seed_gstr1_recon_org(db_session)
    party_id = _party(db_session, org_id, state_code="MH", gstin=_GSTIN)
    inv = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2024, 8, 25),
    )
    _cancel(
        db_session, org_id=org_id, invoice=inv, at=datetime.datetime(2024, 9, 10, 6, tzinfo=_UTC)
    )
    sep = reports_service.compute_gstr1(
        db_session, org_id=org_id, firm_id=firm_id, period="2024-09", can_view_pii=False
    )
    assert sep.cdnr[0].gstin == "************1Z5"


# ──────────────────────────────────────────────────────────────────────
# 3. Cross-period B2CL / export → CDNUR
# ──────────────────────────────────────────────────────────────────────


def test_cross_period_b2cl_cancel_reports_cdnur(db_session: OrmSession) -> None:
    org_id, firm_id, item_id = _seed_gstr1_recon_org(db_session)  # firm MH
    party_id = _party(db_session, org_id, state_code="GJ")
    inv = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2024, 8, 20),
        price="200000",
        ship_to_state="GJ",
    )
    _cancel(
        db_session, org_id=org_id, invoice=inv, at=datetime.datetime(2024, 9, 10, 6, tzinfo=_UTC)
    )

    aug = _assert_books_equal_return(db_session, org_id=org_id, firm_id=firm_id, period="2024-08")
    assert len(aug.b2cl) == 1 and aug.b2cl[0].igst == Decimal("10000.00")

    sep = _assert_books_equal_return(db_session, org_id=org_id, firm_id=firm_id, period="2024-09")
    assert sep.b2cl == [] and sep.cdnr == []
    assert len(sep.cdnur) == 1
    note = sep.cdnur[0]
    assert note.ur_type == "B2CL"
    assert note.note_type == "C"
    assert note.invoice_date == datetime.date(2024, 8, 20)
    assert note.place_of_supply_state == "GJ"
    assert note.note_value == Decimal("210000.00")
    assert (note.taxable_value, note.igst, note.cgst, note.sgst) == (
        Decimal("200000.00"),
        Decimal("10000.00"),
        Decimal("0"),
        Decimal("0"),
    )


def test_cross_period_export_cancel_reports_cdnur_expwop(db_session: OrmSession) -> None:
    org_id, firm_id, item_id = _seed_gstr1_recon_org(db_session)
    party_id = _party(db_session, org_id, state_code=None, is_export=True)
    inv = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2024, 8, 20),
        gst_rate="0",
    )
    _cancel(
        db_session, org_id=org_id, invoice=inv, at=datetime.datetime(2024, 9, 10, 6, tzinfo=_UTC)
    )
    aug = _assert_books_equal_return(db_session, org_id=org_id, firm_id=firm_id, period="2024-08")
    assert len(aug.export) == 1
    sep = _assert_books_equal_return(db_session, org_id=org_id, firm_id=firm_id, period="2024-09")
    assert len(sep.cdnur) == 1
    assert sep.cdnur[0].ur_type == "EXPWOP"
    assert sep.cdnur[0].taxable_value == Decimal("1000.00")


# ──────────────────────────────────────────────────────────────────────
# 4. Cross-period B2CS → net off B2CS + HSN of the cancel month
# ──────────────────────────────────────────────────────────────────────


def test_cross_period_b2cs_cancel_nets_off_cancel_month_b2cs_and_hsn(
    db_session: OrmSession,
) -> None:
    org_id, firm_id, item_id = _seed_gstr1_recon_org(db_session)
    party_id = _party(db_session, org_id, state_code="MH")
    aug_inv = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2024, 8, 20),
        ship_to_state="MH",
    )
    # A September B2CS sale in the same state + rate (3 x ₹1,000 @5%).
    _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2024, 9, 3),
        qty="3",
        ship_to_state="MH",
    )
    _cancel(
        db_session,
        org_id=org_id,
        invoice=aug_inv,
        at=datetime.datetime(2024, 9, 10, 6, tzinfo=_UTC),
    )

    aug = _assert_books_equal_return(db_session, org_id=org_id, firm_id=firm_id, period="2024-08")
    assert len(aug.b2cs) == 1 and aug.b2cs[0].taxable_value == Decimal("1000.00")

    sep = _assert_books_equal_return(db_session, org_id=org_id, firm_id=firm_id, period="2024-09")
    assert sep.cdnr == [] and sep.cdnur == []
    assert len(sep.b2cs) == 1
    b2cs = sep.b2cs[0]
    assert (b2cs.place_of_supply_state, b2cs.gst_rate) == ("MH", Decimal("5"))
    assert (b2cs.taxable_value, b2cs.cgst, b2cs.sgst, b2cs.igst) == (
        Decimal("2000.00"),
        Decimal("50.00"),
        Decimal("50.00"),
        Decimal("0"),
    )
    assert b2cs.invoice_count == 1  # the credit note is not an invoice
    assert len(sep.hsn) == 1
    hsn = sep.hsn[0]
    assert (hsn.total_qty, hsn.taxable_value, hsn.cgst, hsn.sgst, hsn.total_value) == (
        Decimal("2"),
        Decimal("2000.00"),
        Decimal("50.00"),
        Decimal("50.00"),
        Decimal("2100.00"),
    )


def test_cross_period_b2cs_cancel_without_cancel_month_sales_emits_negative_row(
    db_session: OrmSession,
) -> None:
    org_id, firm_id, item_id = _seed_gstr1_recon_org(db_session)
    party_id = _party(db_session, org_id, state_code="MH")
    inv = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2024, 8, 20),
        ship_to_state="MH",
    )
    _cancel(
        db_session, org_id=org_id, invoice=inv, at=datetime.datetime(2024, 9, 10, 6, tzinfo=_UTC)
    )
    sep = _assert_books_equal_return(db_session, org_id=org_id, firm_id=firm_id, period="2024-09")
    assert len(sep.b2cs) == 1
    row = sep.b2cs[0]
    assert (row.place_of_supply_state, row.gst_rate, row.invoice_count) == ("MH", Decimal("5"), 0)
    assert (row.taxable_value, row.cgst, row.sgst) == (
        Decimal("-1000.00"),
        Decimal("-25.00"),
        Decimal("-25.00"),
    )


# ──────────────────────────────────────────────────────────────────────
# 5. IST boundary
# ──────────────────────────────────────────────────────────────────────


def test_ist_boundary_cancel_is_cross_period(db_session: OrmSession) -> None:
    """Invoice 31-Aug; cancel at 2024-08-31T19:00Z = 01-Sep 00:30 IST →
    cross-period (September credit note)."""
    org_id, firm_id, item_id = _seed_gstr1_recon_org(db_session)
    party_id = _party(db_session, org_id, state_code="MH", gstin=_GSTIN)
    inv = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2024, 8, 31),
    )
    _cancel(
        db_session, org_id=org_id, invoice=inv, at=datetime.datetime(2024, 8, 31, 19, tzinfo=_UTC)
    )
    cn = _credit_note(db_session, org_id=org_id, invoice=inv)
    assert cn.voucher_date == datetime.date(2024, 9, 1)

    aug = _assert_books_equal_return(db_session, org_id=org_id, firm_id=firm_id, period="2024-08")
    assert len(aug.b2b) == 1 and aug.cdnr == []
    sep = _assert_books_equal_return(db_session, org_id=org_id, firm_id=firm_id, period="2024-09")
    assert len(sep.cdnr) == 1 and sep.cdnr[0].note_date == datetime.date(2024, 9, 1)


# ──────────────────────────────────────────────────────────────────────
# 6. §34 time limit
# ──────────────────────────────────────────────────────────────────────


def _may_2024_invoice(db_session: OrmSession) -> tuple[uuid.UUID, SalesInvoice]:
    org_id, firm_id, item_id = _seed_gstr1_recon_org(db_session)
    party_id = _party(db_session, org_id, state_code="MH", gstin=_GSTIN)
    inv = _invoice(
        db_session,
        org_id=org_id,
        firm_id=firm_id,
        party_id=party_id,
        item_id=item_id,
        invoice_date=datetime.date(2024, 5, 15),
    )
    return org_id, inv


def test_credit_note_allowed_on_30_nov_after_fy_end(db_session: OrmSession) -> None:
    org_id, inv = _may_2024_invoice(db_session)
    # 2025-11-30T18:00Z = 30-Nov 23:30 IST — last permitted day.
    cancelled = _cancel(
        db_session, org_id=org_id, invoice=inv, at=datetime.datetime(2025, 11, 30, 18, tzinfo=_UTC)
    )
    assert cancelled.lifecycle_status == InvoiceLifecycleStatus.CANCELLED
    assert _credit_note(db_session, org_id=org_id, invoice=inv).voucher_date == datetime.date(
        2025, 11, 30
    )


def test_credit_note_refused_after_30_nov_following_fy_end(db_session: OrmSession) -> None:
    org_id, inv = _may_2024_invoice(db_session)
    # 2025-11-30T18:31Z = 01-Dec 00:01 IST — past the §34 time limit.
    with pytest.raises(InvoiceStateError) as exc:
        _cancel(
            db_session,
            org_id=org_id,
            invoice=inv,
            at=datetime.datetime(2025, 11, 30, 18, 31, tzinfo=_UTC),
        )
    msg = str(exc.value)
    assert "34" in msg and "30-Nov-2025" in msg and "CA" in msg
    # Nothing was posted or changed.
    db_session.refresh(inv)
    assert inv.lifecycle_status != InvoiceLifecycleStatus.CANCELLED
    n_cn = db_session.execute(
        select(func.count())
        .select_from(Voucher)
        .where(Voucher.org_id == org_id, Voucher.voucher_type == VoucherType.CREDIT_NOTE)
    ).scalar_one()
    assert n_cn == 0


@pytest.mark.parametrize(
    ("invoice_date", "deadline"),
    [
        (datetime.date(2024, 4, 1), datetime.date(2025, 11, 30)),  # FY 2024-25 start
        (datetime.date(2024, 5, 15), datetime.date(2025, 11, 30)),
        (datetime.date(2025, 3, 31), datetime.date(2025, 11, 30)),  # FY 2024-25 end
        (datetime.date(2025, 1, 10), datetime.date(2025, 11, 30)),
        (datetime.date(2025, 4, 1), datetime.date(2026, 11, 30)),  # FY 2025-26
    ],
)
def test_credit_note_deadline_is_30_nov_after_fy_end(
    invoice_date: datetime.date, deadline: datetime.date
) -> None:
    from app.service import gst_service

    assert gst_service.credit_note_deadline(invoice_date) == deadline


def test_gst_local_date_uses_ist() -> None:
    from app.service import gst_service

    assert gst_service.gst_local_date(
        datetime.datetime(2024, 8, 31, 18, 29, tzinfo=_UTC)
    ) == datetime.date(2024, 8, 31)
    assert gst_service.gst_local_date(
        datetime.datetime(2024, 8, 31, 18, 30, tzinfo=_UTC)
    ) == datetime.date(2024, 9, 1)


# ──────────────────────────────────────────────────────────────────────
# HTTP: GET /reports/gstr1 carries cdnr / cdnur (JSON + XLSX); the §34
# refusal surfaces as a 409 on POST /invoices/{id}/cancel.
# ──────────────────────────────────────────────────────────────────────


def test_gstr1_endpoint_returns_cdnr_and_xlsx_sheet(
    http_client: TestClient, sync_engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    me = _signup_owner(http_client)
    org_id = uuid.UUID(me["org_id"])
    party_id = _seed_b2b_party(sync_engine, org_id=org_id)
    item_id = _seed_item(sync_engine, org_id=org_id)
    invoice_id = _create_and_finalize_invoice(
        http_client, me, party_id=party_id, item_id=item_id, invoice_date="2024-08-25"
    )
    monkeypatch.setattr(
        sales_service, "_utcnow", lambda: datetime.datetime(2024, 9, 10, 6, tzinfo=_UTC)
    )
    resp = http_client.post(
        f"/invoices/{invoice_id}/cancel",
        headers=_auth(me["access_token"]),
        json={"reason": "goods rejected"},
    )
    assert resp.status_code == 200, resp.text

    aug = http_client.get("/reports/gstr1?period=2024-08", headers=_auth(me["access_token"]))
    assert aug.status_code == 200, aug.text
    assert [r["sales_invoice_id"] for r in aug.json()["b2b"]] == [invoice_id]
    assert aug.json()["cdnr"] == [] and aug.json()["cdnur"] == []

    sep = http_client.get("/reports/gstr1?period=2024-09", headers=_auth(me["access_token"]))
    assert sep.status_code == 200, sep.text
    body = sep.json()
    assert body["b2b"] == [] and body["cdnur"] == []
    assert len(body["cdnr"]) == 1
    note = body["cdnr"][0]
    assert note["sales_invoice_id"] == invoice_id
    assert note["note_type"] == "C"
    assert note["note_date"] == "2024-09-10"
    assert note["invoice_date"] == "2024-08-25"
    assert note["gstin"] is not None and note["gstin"].endswith("1Z5")
    assert "ur_type" not in note
    assert Decimal(note["taxable_value"]) == Decimal("1000.00")
    assert Decimal(note["cgst"]) == Decimal("25.00") == Decimal(note["sgst"])

    xlsx = http_client.get(
        "/reports/gstr1?period=2024-09&format=xlsx", headers=_auth(me["access_token"])
    )
    assert xlsx.status_code == 200, xlsx.text
    wb = load_workbook(io.BytesIO(xlsx.content))
    ws = wb["CDNR"]
    header = [c.value for c in ws[1]]
    values = [c.value for c in ws[2]]
    row = dict(zip(header, values, strict=True))
    assert row["Note #"] == note["note_number"]
    assert row["Original invoice #"] == note["invoice_number"]
    assert row["Note type"] == "C"
    assert Decimal(str(row["CGST"])) == Decimal("25.00")


def test_cancel_endpoint_refuses_after_s34_time_limit(
    http_client: TestClient, sync_engine: Engine, monkeypatch: pytest.MonkeyPatch
) -> None:
    me = _signup_owner(http_client)
    org_id = uuid.UUID(me["org_id"])
    party_id = _seed_b2b_party(sync_engine, org_id=org_id)
    item_id = _seed_item(sync_engine, org_id=org_id)
    invoice_id = _create_and_finalize_invoice(
        http_client, me, party_id=party_id, item_id=item_id, invoice_date="2024-05-15"
    )
    monkeypatch.setattr(
        sales_service, "_utcnow", lambda: datetime.datetime(2025, 12, 1, 6, tzinfo=_UTC)
    )
    resp = http_client.post(
        f"/invoices/{invoice_id}/cancel",
        headers=_auth(me["access_token"]),
        json={"reason": "late"},
    )
    assert resp.status_code == 409, resp.text
    assert "34" in resp.text and "CA" in resp.text
